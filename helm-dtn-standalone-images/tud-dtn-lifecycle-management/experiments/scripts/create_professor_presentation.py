#!/usr/bin/env python3
"""Build an editable professor-facing PowerPoint from real experiment artifacts."""

from pathlib import Path
from datetime import datetime
import re

import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE


ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "store-forward" / "results"
CSV = RESULTS / "persistent_store_forward_results.csv"
GRAPHS = RESULTS / "graphs"
OUT = RESULTS / "Helm_DTN_OCI_Image_Transfer_Experiments.pptx"

NAVY = RGBColor(21, 59, 91)
BLUE = RGBColor(0, 114, 178)
GREEN = RGBColor(0, 158, 115)
ORANGE = RGBColor(230, 159, 0)
RED = RGBColor(213, 94, 0)
PURPLE = RGBColor(122, 81, 149)
PINK = RGBColor(204, 121, 167)
SLATE = RGBColor(71, 85, 105)
LIGHT = RGBColor(241, 245, 249)
WHITE = RGBColor(255, 255, 255)


def truth(series):
    return series.astype(str).str.lower().eq("true")


def expected(run_id):
    match = re.search(r"_persistent_sf_(\d+)s_", str(run_id))
    return int(match.group(1)) if match else None


df = pd.read_csv(CSV)
df["scenario"] = df.network_scenario.astype(str).str.lower()
df["outage"] = pd.to_numeric(df.interruption_duration_seconds, errors="coerce")
df["expected"] = df.run_id.map(expected)
scenario_df = df[df.scenario.isin(["leo", "lunar", "mars"]) & (df.outage == df.expected)].copy()
scenario_df["valid_success"] = truth(scenario_df.success) & truth(scenario_df.digest_match)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def background(slide, color=WHITE):
    fill = slide.background.fill
    fill.solid(); fill.fore_color.rgb = color


def title(slide, text, subtitle=None):
    box = slide.shapes.add_textbox(Inches(.55), Inches(.28), Inches(12.2), Inches(.65))
    p = box.text_frame.paragraphs[0]
    p.text = text; p.font.name = "Aptos Display"; p.font.size = Pt(27)
    p.font.bold = True; p.font.color.rgb = NAVY
    if subtitle:
        sub = slide.shapes.add_textbox(Inches(.58), Inches(.90), Inches(12), Inches(.38))
        q = sub.text_frame.paragraphs[0]; q.text = subtitle
        q.font.name = "Aptos"; q.font.size = Pt(11); q.font.color.rgb = SLATE
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.55), Inches(1.22), Inches(12.2), Inches(.035))
    line.fill.solid(); line.fill.fore_color.rgb = BLUE; line.line.fill.background()


def footer(slide, n):
    box = slide.shapes.add_textbox(Inches(.55), Inches(7.14), Inches(12.2), Inches(.2))
    p = box.text_frame.paragraphs[0]
    p.text = f"Helm-DTN / µD3TN OCI image-transfer evaluation  •  live snapshot {datetime.now():%Y-%m-%d}  •  {n}"
    p.font.size = Pt(7.5); p.font.color.rgb = SLATE; p.alignment = PP_ALIGN.RIGHT


def add_slide(name, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); background(slide); title(slide, name, subtitle)
    footer(slide, len(prs.slides)); return slide


def textbox(slide, x, y, w, h, text, size=17, color=SLATE, bold=False,
            fill=None, border=None, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        shape.fill.solid(); shape.fill.fore_color.rgb = fill
    else:
        shape.fill.background()
    if border:
        shape.line.color.rgb = border
    else:
        shape.line.fill.background()
    tf = shape.text_frame; tf.word_wrap = True; tf.vertical_anchor = valign
    tf.margin_left = tf.margin_right = Inches(.12)
    p = tf.paragraphs[0]; p.text = text; p.alignment = align
    p.font.name = "Aptos"; p.font.size = Pt(size); p.font.color.rgb = color; p.font.bold = bold
    return shape


def bullets(slide, items, x=.7, y=1.5, w=12, h=5.2, size=20, color=SLATE):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame; tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item; p.level = 0; p.font.name = "Aptos"; p.font.size = Pt(size)
        p.font.color.rgb = color; p.space_after = Pt(10); p.text = "• " + p.text
    return shape


def picture(slide, path, x, y, w=None, h=None):
    kwargs = {}
    if w: kwargs["width"] = Inches(w)
    if h: kwargs["height"] = Inches(h)
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), **kwargs)


def table(slide, rows, cols, data, x, y, w, h, widths=None, font=10):
    shape = slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(h))
    tab = shape.table
    if widths:
        for i, value in enumerate(widths): tab.columns[i].width = Inches(value)
    for r in range(rows):
        for c in range(cols):
            cell = tab.cell(r, c); cell.text = str(data[r][c])
            cell.margin_left = cell.margin_right = Inches(.05)
            cell.fill.solid(); cell.fill.fore_color.rgb = NAVY if r == 0 else (LIGHT if r % 2 else WHITE)
            for p in cell.text_frame.paragraphs:
                p.font.name = "Aptos"; p.font.size = Pt(font)
                p.font.bold = r == 0; p.font.color.rgb = WHITE if r == 0 else SLATE
                p.alignment = PP_ALIGN.CENTER
    return tab


# 1 — Title
slide = prs.slides.add_slide(prs.slide_layouts[6]); background(slide, NAVY)
textbox(slide, .75, 1.05, 11.8, 1.5, "Transferring Talos/Kubernetes OCI Images over Helm-DTN and µD3TN",
        32, WHITE, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
textbox(slide, 1.25, 2.7, 10.8, .8, "Evaluation under LEO, Lunar and Mars contact conditions",
        20, RGBColor(191, 219, 254), False, align=PP_ALIGN.CENTER)
textbox(slide, 2.1, 4.05, 9.1, 1.15,
        "OCI packaging • AAP2/BPv7 • TCPCL • persistent store-and-forward • digest verification",
        17, WHITE, False, fill=RGBColor(30, 64, 96), border=BLUE, align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE)
textbox(slide, 3.7, 6.25, 6, .35, "Anwesha — experiment progress presentation", 12,
        RGBColor(203, 213, 225), align=PP_ALIGN.CENTER)

# 2 — Motivation
slide = add_slide("Experiment motivation", "Why image delivery is required before the Space-side cluster can use its dependencies")
bullets(slide, [
    "A disconnected Space-side environment cannot depend on continuous access to public registries.",
    "Talos/Kubernetes nodes require networking, DNS, control-plane, node-agent and sandbox images.",
    "Long-delay or intermittent links break assumptions made by ordinary always-connected registry workflows.",
    "The research question: can the existing OCI/AAP2/µD3TN route preserve and deliver an image after contact returns?",
], .8, 1.55, 11.7, 4.2, 20)
textbox(slide, 1.15, 5.75, 11.0, .72,
        "Goal: evaluate the existing transfer implementation — not build a Kubernetes cluster and not rewrite the sender/receiver.",
        17, NAVY, True, fill=RGBColor(239, 246, 255), border=BLUE, align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE)

# 3 — Scope/requirements
slide = add_slide("Scope and requirements", "What was evaluated and what was deliberately excluded")
table(slide, 6, 2, [
    ["In scope", "Out of scope"],
    ["Identify required cluster images", "Create/install the Kubernetes cluster"],
    ["Transfer through the existing route", "Rewrite OCI packaging or AAP2"],
    ["10 repetitions per condition", "Invent or repair missing measurements"],
    ["CPU, RSS, payload, duration, throughput", "Claim local results as spacecraft performance"],
    ["Normal/contact-loss recovery + digests", "Claim Mars propagation unless delay was applied"],
], .9, 1.55, 11.5, 4.9, [5.75, 5.75], 14)

# 4 — Why Helm-DTN
slide = add_slide("Why the existing Helm-DTN route?", "Helm deploys/configures the environment; it does not copy the image")
bullets(slide, [
    "Reuses the project’s already implemented sender, receiver, OCI layout, AAP2 and digest-verification path.",
    "Provides a repeatable µD3TN deployment/routing environment instead of a one-off transport setup.",
    "Uses BPv7 bundles so application delivery is separated from one continuous end-to-end connection.",
    "Keeps registry compatibility: the receiver reconstructs OCI content and pushes it into the Space-side registry.",
    "Avoids adding another external image-transfer library; evaluation remains focused on the existing design.",
], .8, 1.55, 11.7, 4.9, 18)

# 5 — Architecture graphic
slide = add_slide("End-to-end transfer architecture", "One OCI TAR payload is submitted as one AAP2 ADU")
picture(slide, GRAPHS / "method_explanation_chart.png", .55, 1.35, w=12.25)

# 6 — Images
slide = add_slide("Image sources and selection", "Nine Kubernetes runtime images plus one Talos installer image")
image_rows = [
    ["Image", "Source registry", "Purpose"],
    ["flannel:v0.27.4", "ghcr.io/siderolabs", "Pod CNI/networking"],
    ["coredns:v1.13.2", "registry.k8s.io", "Cluster DNS"],
    ["etcd:v3.6.7", "registry.k8s.io", "Cluster state"],
    ["kube-apiserver:v1.35.0", "registry.k8s.io", "Kubernetes API"],
    ["controller-manager:v1.35.0", "registry.k8s.io", "Controllers"],
    ["scheduler:v1.35.0", "registry.k8s.io", "Pod scheduling"],
    ["kube-proxy:v1.35.0", "registry.k8s.io", "Service networking"],
    ["kubelet:v1.35.0", "ghcr.io/siderolabs", "Node agent"],
    ["pause:3.10", "registry.k8s.io", "Pod sandbox"],
    ["installer:v1.12.4", "ghcr.io/siderolabs", "Talos install/upgrade"],
]
table(slide, len(image_rows), 3, image_rows, .75, 1.42, 11.8, 5.55, [3.6, 3.3, 4.9], 10.5)

# 7 — Method
slide = add_slide("Experimental method and evidence", "The same measurement contract is used for every run")
bullets(slide, [
    "Pull image → OCI layout → TAR + metadata → one AAP2 ADU.",
    "Temporarily interrupt only the Earth-to-Space µD3TN/TCPCL path.",
    "Keep sender/receiver alive; restore contact and observe delayed delivery.",
    "Push to the Space registry and compare source/destination SHA-256 digests.",
    "Record payload, timestamps, duration, CPU, RSS, throughput, interruption, reconnection, success and error.",
    "Store one detailed log per run and one row per attempt in the combined CSV.",
], .75, 1.48, 12, 5.2, 18)

# 8 — Scenarios
slide = add_slide("Why LEO, Lunar and Mars?", "Controlled profiles represent increasingly difficult communication conditions")
table(slide, 4, 5, [
    ["Scenario", "Bandwidth", "Applied RTT", "Reference RTT", "Interpretation"],
    ["LEO", "15 Mbit/s", "50 ms", "50 ms", "Low-delay orbital profile"],
    ["Lunar", "20 Mbit/s", "2.6 s", "2.6 s", "1.3 s one-way delay"],
    ["Mars", "5 Mbit/s", "0 ms", "1800 s", "Contact-only; Mars delay not emulated"],
], .65, 1.5, 12.0, 2.4, [1.5, 1.7, 1.7, 1.8, 5.3], 12)
textbox(slide, .8, 4.25, 5.8, 1.45,
        "Bandwidth ≠ RTT\nCapacity depends on link budget, power, antennas and coding. RTT mainly reflects propagation distance.",
        16, NAVY, True, fill=RGBColor(239, 246, 255), border=BLUE, align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE)
textbox(slide, 6.75, 4.25, 5.75, 1.45,
        "Mars limitation\n5 Mbit/s + contact disruption were applied. A literal 900 s one-way delay was not.",
        16, RGBColor(131, 24, 67), True, fill=RGBColor(253, 242, 248), border=PINK,
        align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)

# 9 — RTT graphic
slide = add_slide("Reference RTT versus RTT actually applied", "A graph must never substitute a reference value for an applied measurement")
picture(slide, GRAPHS / "corrected_timelines" / "rtt_reference_vs_applied.png", 1.35, 1.35, w=10.65)

# 10 — Contact window
slide = add_slide("Contact-window experiment", "Same contact logic for LEO, Lunar and Mars; outage = 10 s, 5 min or 1 h")
textbox(slide, .8, 1.55, 2.25, .65, "CONTACT ON", 18, WHITE, True, fill=GREEN,
        align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
textbox(slide, 3.05, 1.55, 7.0, .65, "CONTACT OFF — bundle retained at Earth", 18, NAVY, True,
        fill=RGBColor(254, 226, 226), border=RED, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
textbox(slide, 10.05, 1.55, 2.25, .65, "RESTORED", 18, WHITE, True, fill=GREEN,
        align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
for y, label, color in [(2.75, "LEO", BLUE), (3.65, "LUNAR", ORANGE), (4.55, "MARS contact-only", PINK)]:
    textbox(slide, .85, y, 1.65, .48, label, 13, NAVY, True, align=PP_ALIGN.RIGHT)
    textbox(slide, 2.7, y, 1.7, .48, "ADU submitted", 12, WHITE, True, fill=color,
            align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    textbox(slide, 4.4, y, 4.7, .48, "10 s / 300 s / 3600 s interruption", 12, NAVY, True,
            fill=RGBColor(255, 237, 213), border=RED, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    textbox(slide, 9.1, y, 2.5, .48, "TCPCL + delivery", 12, WHITE, True, fill=color,
            align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
textbox(slide, 1.1, 5.65, 11.1, .75,
        "This is a controlled experimental contact window—not a real satellite pass, lunar visibility schedule or DSN Mars schedule.",
        15, NAVY, True, fill=LIGHT, border=RGBColor(203, 213, 225), align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE)

# 11 — Current coverage
slide = add_slide("Current result coverage", "Live matrix snapshot; unfinished images must not be presented as completed")
summary = scenario_df.groupby("image_name").agg(
    attempts=("run_id", "size"), successes=("valid_success", "sum")
).reset_index()
coverage = [["Image", "Recorded attempts", "Digest-verified successes", "Status"]]
configured = ["flannel", "coredns", "etcd", "kube-apiserver", "kube-controller-manager",
              "kube-scheduler", "kube-proxy", "kubelet", "pause", "talos-installer"]
for name in configured:
    row = summary[summary.image_name == name]
    attempts = int(row.attempts.iloc[0]) if len(row) else 0
    successes = int(row.successes.iloc[0]) if len(row) else 0
    status = "complete/mostly complete" if attempts >= 90 else ("in progress" if attempts else "not reached")
    coverage.append([name, attempts, successes, status])
table(slide, len(coverage), 4, coverage, .8, 1.42, 11.7, 5.55, [3.3, 2.3, 3.0, 3.1], 10.5)

# 12 — Performance chart
slide = add_slide("What the measured trends show", "Duration rises and effective throughput falls as the outage becomes longer")
picture(slide, GRAPHS / "scenario_performance_comparison.png", .35, 1.4, w=12.65)

# 13 — Why results look like this
slide = add_slide("Why do the curves rise and fall?", "Interruption waiting time dominates the complete wall-clock measurement")
table(slide, 5, 3, [
    ["Observation", "Cause", "Correct interpretation"],
    ["Duration rises", "Bundle waits for contact restoration", "Expected DTN delay, plus recovery overhead"],
    ["Throughput falls", "Payload ÷ total wall-clock duration", "Not necessarily lower active-link capacity"],
    ["High error bars", "WSL suspension and variable reconnection", "Operational variability, not image corruption"],
    ["Digests match", "Receiver verifies/pushes original OCI content", "Integrity succeeded after delivery"],
], .7, 1.55, 12, 4.6, [3.0, 4.1, 4.9], 13)
textbox(slide, 1.0, 6.25, 11.3, .55,
        "Longer total time can coexist with a successful digest: DTN optimizes eventual delivery, not low latency.",
        16, NAVY, True, fill=RGBColor(240, 253, 250), border=GREEN, align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE)

# 14 — TCPCL
slide = add_slide("TCPCL: when it is good—and when it is not", "TCPCL is the transport convergence layer, not the persistent DTN store")
table(slide, 5, 2, [
    ["Works well when…", "Becomes difficult when…"],
    ["An IP path currently exists", "The contact disappears for a long time"],
    ["RTT is within TCP/handshake timeout limits", "RTT is extreme, e.g. literal Mars delay"],
    ["The session remains stable", "Host/network suspension breaks timers and connections"],
    ["The link becomes active before recall", "A stored bundle is recalled before handshake completion"],
], .75, 1.55, 11.8, 3.8, [5.9, 5.9], 14)
textbox(slide, 1.0, 5.75, 11.3, .8,
        "Observed fix: configure contact → wait 15 s for TCPCL handshake → recall the SQLite bundle.",
        18, NAVY, True, fill=RGBColor(255, 247, 237), border=ORANGE, align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE)

# 15 — DTN store-forward learning
slide = add_slide("What persistent DTN changed", "The original ADU survives contact loss without restarting the sender")
steps = [
    ("1", "Sender submits once", BLUE), ("2", "Earth stores bundle", ORANGE),
    ("3", "Contact unavailable", RED), ("4", "TCPCL reconnects", PURPLE),
    ("5", "Stored bundle recalled", GREEN), ("6", "Registry + digest verified", BLUE),
]
for i, (num, text, color) in enumerate(steps):
    x = .75 + (i % 3) * 4.15; y = 1.65 + (i // 3) * 2.25
    textbox(slide, x, y, .7, .7, num, 22, WHITE, True, fill=color,
            align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    textbox(slide, x + .82, y, 3.0, .7, text, 16, NAVY, True, fill=LIGHT,
            border=color, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
textbox(slide, 1.0, 6.0, 11.3, .65,
        "Learning: DTN store-and-forward provides eventual delivery across disruption; TCPCL alone does not provide custody/persistence.",
        16, NAVY, True, fill=RGBColor(240, 253, 250), border=GREEN, align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE)

# 16 — Limitations and next steps
slide = add_slide("Limitations, requirements and next steps", "What can be claimed now—and what requires a separate experiment")
table(slide, 7, 2, [
    ["Current limitation", "Requirement / next step"],
    ["WSL pauses sleep timers", "Use an always-on Linux host or absolute wall-clock scheduler"],
    ["Mars applied RTT = 0 ms", "Tune timeouts and run a separate 900 s one-way-delay experiment"],
    ["Local CPU/RSS values", "Repeat on representative edge/space hardware"],
    ["One large ADU", "Evaluate fragmentation/chunking and partial resume"],
    ["Asymmetric links not modeled", "Apply distinct forward/return bandwidth profiles"],
    ["Matrix still running", "Finish ten images, regenerate summaries, freeze final dataset"],
], .7, 1.45, 12, 4.85, [5.5, 6.5], 13)
textbox(slide, 1.0, 6.4, 11.3, .48,
        "Do not claim physical Mars propagation performance from the current contact-only Mars results.",
        15, RGBColor(131, 24, 67), True, fill=RGBColor(253, 242, 248), border=PINK,
        align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)

# 17 — Conclusion
slide = add_slide("Conclusion", "What the experiment demonstrates so far")
bullets(slide, [
    "The existing Helm-DTN/µD3TN OCI route can deliver cluster images and preserve digest integrity.",
    "Short and medium outages usually recover through TCPCL reconnection.",
    "Long outages require persistent Earth-side storage and handshake-aware bundle recall.",
    "LEO and Lunar results include applied propagation RTT; Mars results are contact/bandwidth recovery only.",
    "The most important operational issue is WSL suspension extending configured one-hour interruptions.",
], .8, 1.55, 11.7, 4.4, 19)
textbox(slide, 1.0, 5.95, 11.3, .8,
        "Core message: the system prioritizes eventual, integrity-verified delivery across disrupted contact—not minimum latency.",
        18, NAVY, True, fill=RGBColor(239, 246, 255), border=BLUE, align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE)

# 18 — Evidence/references
slide = add_slide("Evidence and references", "Repository artifacts used to build this presentation")
bullets(slide, [
    "Raw results: experiments/store-forward/results/persistent_store_forward_results.csv",
    "Detailed logs: experiments/store-forward/logs/{leo,lunar,mars}/",
    "Network profiles: experiments/config/network_scenarios.yaml",
    "Experiment scripts: run_full_persistent_matrix.sh and run_persistent_store_forward.sh",
    "Architecture/method report: experiments/store-forward/PROFESSOR_EXPLANATION_CHART.md",
    "NASA NESC Mars communications report: Mars-to-Earth rates approximately 1–5 Mbit/s",
    "JPL Mars optical-link study: modeled downlink rates approximately 5–40 Mbit/s",
    "Professor-provided IMC paper: interruption methodology/reference context",
], .75, 1.45, 12, 5.7, 15)

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(f"Created {OUT} with {len(prs.slides)} slides from {len(scenario_df)} scenario-labelled attempts")
